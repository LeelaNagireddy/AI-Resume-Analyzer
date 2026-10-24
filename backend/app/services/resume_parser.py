import pymupdf
from docx import Document
import io


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF resume."""

    document = pymupdf.open(file_path)

    text = ""

    for page in document:
        text += page.get_text()

    document.close()

    return text.strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX resume."""

    document = Document(file_path)

    text = "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
    )

    return text.strip()


def extract_resume_text(file_path: str) -> str:
    """Extract text based on the resume file extension."""

    if file_path.lower().endswith(".pdf"):
        return extract_text_from_pdf(file_path)

    if file_path.lower().endswith(".docx"):
        return extract_text_from_docx(file_path)

    raise ValueError("Unsupported file format. Use PDF or DOCX.")


def extract_text_from_pdf_bytes(data: bytes) -> str:
    """Extract text from PDF bytes (nothing is written to disk)."""

    with pymupdf.open(stream=data, filetype="pdf") as document:
        text = "".join(page.get_text() for page in document)

    return text.strip()


def extract_text_from_docx_bytes(data: bytes) -> str:
    """Extract text from DOCX bytes (nothing is written to disk)."""

    document = Document(io.BytesIO(data))

    text = "\n".join(paragraph.text for paragraph in document.paragraphs)

    return text.strip()


def extract_resume_text_from_bytes(data: bytes, extension: str) -> str:
    """Extract resume text from uploaded bytes, based on the file extension."""

    if extension == ".pdf":
        return extract_text_from_pdf_bytes(data)

    if extension == ".docx":
        return extract_text_from_docx_bytes(data)

    raise ValueError("Unsupported file format. Use PDF or DOCX.")